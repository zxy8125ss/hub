"""根据 status.json 生成「App总览.xlsx」。由状态检查任务每 30 分钟运行一次，结果放在 status 分支，总控台页面提供下载。
用法：python3 scripts/excel.py out/status.json out/App总览.xlsx"""
import datetime as dt, json, sys, re
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BJ = dt.timezone(dt.timedelta(hours=8))
BAD = {'failure', 'timed_out', 'startup_failure', 'action_required'}
IGNORE = {'cancelled', 'skipped', 'neutral', 'stale'}
JOBNAME = {'pages build and deployment': '更新网页', 'Deploy to gh-pages': '更新网页', '部署': '更新网页',
           '部署到 GitHub Pages': '更新网页', 'Deploy to GitHub Pages': '更新网页', 'pages': '更新网页'}
KIND = {'local': '只在手机里', 'cloud': '已存 GitHub', 'none': '无需备份', 'minor': '影响不大', 'unknown': '未登记'}
WEEK = '一二三四五六日'

INK, SAND, CARD = '29473A', 'E7DCCB', 'FBF6EE'
FILL = {'正常': 'DCEBDF', '留意': 'F6E5C6', '失败': 'F5D7CF', '无任务': 'E9E1D4'}
FONT = 'Arial'


def P(s):
    return dt.datetime.fromisoformat(s.replace('Z', '+00:00')) if s else None


def bj(s, fmt='%Y-%m-%d %H:%M'):
    t = P(s) if isinstance(s, str) else s
    return t.astimezone(BJ).strftime(fmt) if t else ''


# ---------- cron（UTC）→ 北京时间的说法 ----------
def field(f, lo, hi):
    out = set()
    for part in f.split(','):
        rng, _, step = part.partition('/')
        step = int(step or 1)
        if rng == '*':
            a, b = lo, hi
        elif '-' in rng:
            a, b = map(int, rng.split('-'))
        else:
            a = int(rng); b = hi if step > 1 else a
        out.update(range(a, b + 1, step))
    return {0 if (v == 7 and hi == 6) else v for v in out}


def fires_week(expr):
    p = expr.split()
    if len(p) != 5:
        return {}
    mi, h, d, mo, w = field(p[0], 0, 59), field(p[1], 0, 23), field(p[2], 1, 31), field(p[3], 1, 12), field(p[4], 0, 6)
    dany, wany = p[2] == '*', p[4] == '*'
    start = dt.datetime(2026, 10, 19, tzinfo=dt.timezone.utc)  # 周一
    days = {}
    t = start
    while t < start + dt.timedelta(days=7):
        cron_dow = (t.weekday() + 1) % 7
        ok_d = True if dany and wany else (cron_dow in w if dany else (t.day in d if wany else (t.day in d or cron_dow in w)))
        if ok_d and t.month in mo and t.hour in h and t.minute in mi:
            b = t.astimezone(BJ)
            days.setdefault(b.weekday(), set()).add(b.strftime('%H:%M'))
        t += dt.timedelta(minutes=1)
    return days


def schedule_text(crons):
    days = {}
    for c in crons:
        for k, v in fires_week(c).items():
            days.setdefault(k, set()).update(v)
    if not days:
        return ''
    n = max(len(v) for v in days.values())
    if n > 12:
        return f'每天约 {n} 次（每 {round(24 * 60 / n)} 分钟）'
    groups = {}
    for wd in range(7):
        key = tuple(sorted(days.get(wd, ())))
        if key:
            groups.setdefault(key, []).append(wd)
    parts = []
    for times, wds in groups.items():
        if len(wds) == 7:
            label = '每天'
        elif wds == list(range(wds[0], wds[-1] + 1)) and len(wds) > 1:
            label = f'周{WEEK[wds[0]]}至周{WEEK[wds[-1]]}'
        else:
            label = '、'.join('周' + WEEK[x] for x in wds)
        parts.append(f'{label} {"、".join(times)}')
    return '；'.join(parts)


# ---------- 判断 ----------
def judge(a, now):
    runs = a.get('runs') or []
    wfs = {}
    for r in runs:
        wfs.setdefault(r['wid'], []).append(r)
    jobs, state, why = [], '正常', []
    rank = {'失败': 0, '留意': 1, '正常': 2}
    for rs in wfs.values():
        name = JOBNAME.get(rs[0]['name'], rs[0]['name'])
        done = next((r for r in rs if r['status'] == 'completed' and r['conclusion'] not in IGNORE), None)
        sched = next((r for r in rs if r['event'] == 'schedule'), None)
        crons = (a.get('crons') or {}).get(rs[0]['path'], [])
        st, note = '正常', ''
        if done and done['conclusion'] in BAD:
            st, note = '失败', '最近一次失败'
        elif sched and crons:
            gap = ((a.get('maxGap') or {}).get(name)) or 27
            hrs = (now - P(sched['created_at'])).total_seconds() / 3600
            if hrs > gap:
                st, note = '留意', f'定时任务 {int(hrs)} 小时没跑'
        week = [r for r in rs if r['status'] == 'completed' and r['conclusion'] not in IGNORE and (now - P(r['created_at'])).days < 7]
        ok = [r for r in week if r['conclusion'] not in BAD]
        jobs.append({'name': name, 'sched': schedule_text(crons), 'state': st, 'note': note,
                     'last': rs[0], 'week': len(week), 'ok': len(ok),
                     'fail': (a.get('failSteps') or {}).get(str(done['id'])) if st == '失败' and done else ''})
        if rank[st] < rank[state]:
            state = st
        if note:
            why.append(f'{name}：{note}')
    site = a.get('site') or {}
    if site and not site.get('ok'):
        state = '失败'; why.insert(0, '网页打不开')
    if not runs:
        state = '无任务'
    return state, '；'.join(why), jobs


def deploy_state(a, now):
    runs = a.get('runs') or []
    head = a.get('headHuman') if a.get('deployIgnoreBot') else a.get('head')
    dep = [r for r in runs if JOBNAME.get(r['name']) == '更新网页' or re.search(r'deploy|pages', r.get('path') or '', re.I)]
    if not dep:
        return '没有更新网页的记录', ''
    last = dep[0]
    okr = next((r for r in dep if r['status'] == 'completed' and r['conclusion'] == 'success'), None)
    when = bj(okr['updated_at']) if okr else ''
    if last['status'] != 'completed':
        return '正在更新网页', when
    if last['conclusion'] in BAD:
        return '最近一次更新网页失败', when
    if head and P(last['created_at']) < P(head['date']) - dt.timedelta(minutes=1) and (now - P(head['date'])).total_seconds() > 1800:
        return '有改动还没更新到网页', when
    return '网页已是最新版', when


def version_label(a):
    v = (a.get('version') or {}).get('label')
    if v:
        m = re.match(r'^v?\d{4}[-.]?(\d{2})[-.]?(\d{2})', v)
        return f'{int(m.group(1))}.{m.group(2)} 版' if m else ('v' + v if v[0].isdigit() else v)
    h = a.get('headHuman')
    return f'{bj(h["date"], "%m.%d").lstrip("0")} 版' if h else ''


def size(b):
    if not b:
        return ''
    return f'{max(1, round(b / 1024))} KB' if b < 1048576 else f'{b / 1048576:.1f} MB'


# ---------- 写表 ----------
thin = Side(style='thin', color='D9CDB8')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def sheet(wb, title, headers, rows, widths, first=False, links=()):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c in range(1, len(headers) + 1):
        h = ws.cell(row=1, column=c)
        h.font = Font(name=FONT, bold=True, color='FFFFFF', size=11)
        h.fill = PatternFill('solid', fgColor=INK)
        h.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = BORDER
            if isinstance(cell.value, str) and cell.value in FILL:
                cell.fill = PatternFill('solid', fgColor=FILL[cell.value])
    for col in links:   # 网址列做成可点的链接
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell.value, str) and cell.value.startswith('http'):
                cell.hyperlink = cell.value
                cell.font = Font(name=FONT, size=10, color='1F7A70', underline='single')
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def main(src, dst):
    s = json.load(open(src, encoding='utf-8'))
    now = P(s['generated_at'])
    apps = sorted(s['apps'], key=lambda a: (a['key'] == 'hub', a.get('created_at') or ''))
    wb = Workbook()

    over, upd, jobs_rows, bak = [], [], [], []
    for a in apps:
        state, why, jobs = judge(a, now)
        dep, dep_at = deploy_state(a, now)
        ups = a.get('updates') or []
        b = (a.get('backup') or {}).get('latest') or {}
        ph = a.get('phone') or {}
        auto = [j for j in jobs if j['sched']]
        week = sum(j['week'] for j in jobs); ok = sum(j['ok'] for j in jobs)
        over.append([
            a['title'], a.get('about') or a.get('desc') or '', a['url'], state, why or '—',
            version_label(a), dep, dep_at,
            f"{bj(ups[0]['date'], '%m-%d')} {ups[0]['msg']}" if ups else '',
            a.get('humanWeek') or 0, a.get('autoWeek') or 0,
            '；'.join(f"{j['name']}：{j['sched']}" for j in auto) or '无',
            f'{round(ok / week * 100)}%' if week else '—',
            KIND.get(ph.get('kind'), '未登记'), ph.get('text', ''),
            bj(b.get('at')) if b else '', size(b.get('bytes')) if b else '',
            bj(a.get('created_at'), '%Y-%m-%d'), f"https://github.com/{s.get('owner', 'zxy8125ss')}/{a['repo']}",
        ])
        for u in ups[:5]:
            upd.append([a['title'], bj(u['date']), u['msg']])
        for j in jobs:
            last = j['last']
            res = '运行中' if last['status'] != 'completed' else ('失败' if last['conclusion'] in BAD else ('跳过' if last['conclusion'] in IGNORE else '成功'))
            jobs_rows.append([a['title'], j['name'], j['sched'] or '有改动时', j['state'], f"{j['ok']}/{j['week']}" if j['week'] else '—',
                              bj(last['created_at']), res, j['fail'] or j['note'] or '', last['url']])
        bak.append([a['title'], KIND.get(ph.get('kind'), '未登记'), ph.get('text', ''), bj(b.get('at')) if b else '还没有',
                    size(b.get('bytes')) if b else '', b.get('files', '') if b else '', b.get('url', '') if b else '', a['url']])

    ws = sheet(wb, '总览', ['App', '用途', '网址', '运行状态', '需要留意', '版本', '网页状态', '网页更新于',
                          '最近一次功能更新', '本周功能更新', '本周数据自动更新', '自动任务（北京时间）', '7 天任务成功率',
                          '个人数据存在哪', '怎么备份个人数据', '最近自动备份', '备份大小', '开始做的日期', 'GitHub 仓库'],
               over, [16, 34, 44, 10, 22, 11, 16, 16, 36, 9, 11, 36, 10, 12, 40, 16, 10, 12, 34], first=True, links=(3, 19))
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=1).font = Font(name=FONT, size=11, bold=True, color=INK)
        for c in (10, 11):
            ws.cell(row=r, column=c).alignment = Alignment(horizontal='center', vertical='top')

    sheet(wb, '更新记录', ['App', '时间', '更新内容'], upd, [16, 18, 80])
    sheet(wb, '自动任务', ['App', '任务', '什么时候跑（北京时间）', '状态', '7 天成功/总次数', '最近一次', '最近结果', '说明', '任务记录'],
          jobs_rows, [14, 16, 40, 8, 12, 16, 9, 30, 40], links=(9,))
    sheet(wb, '备份', ['App', '个人数据存在哪', '怎么备份个人数据', '最近自动备份', '大小', '文件数', '下载备份', '打开 App'],
          bak, [14, 12, 46, 16, 9, 8, 46, 32], links=(7, 8))

    info = wb.create_sheet('说明')
    lines = [
        ('这份表是什么', '所有自建 App 的汇总：用途、网址、运行是否正常、版本、自动任务、个人数据和备份。'),
        ('生成时间', bj(s['generated_at']) + '（北京时间）'),
        ('怎么更新', '总控台每 30 分钟自动检查一次并重新生成这份表，新做的 App 发布成网页后会自动加进来。在总控台页面点「下载 App 总览」拿到的总是最新的一份。'),
        ('运行状态', '正常：自动任务都成功、网页能打开；留意：定时任务超时没跑等；失败：最近一次任务失败或网页打不开；无任务：没有自动任务的 App。'),
        ('自动备份', f"每天 03:23 把每个 App 的程序和存在 GitHub 上的数据各打一个压缩包，保留 30 天。最近一次：{bj((s.get('backup') or {}).get('at')) or '还没有'}，共 {size((s.get('backup') or {}).get('total_bytes'))}。"),
        ('个人数据', '「只在手机里」的 App，记录只存在你手机上，自动备份拿不到，需要在 App 里自己导出；导出后可以导入总控台的「我的记录」统一查看和导出。'),
        ('总控台', 'https://zxy8125ss.github.io/hub/'),
    ]
    info.column_dimensions['A'].width = 14
    info.column_dimensions['B'].width = 90
    for i, (k, v) in enumerate(lines, 1):
        info.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, color=INK)
        c = info.cell(row=i, column=2, value=v)
        c.font = Font(name=FONT)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if v.startswith('http'):
            c.hyperlink = v
            c.font = Font(name=FONT, color='1F7A70', underline='single')
    wb.save(dst)
    print(f'已生成 {dst}：{len(apps)} 个 App')


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
